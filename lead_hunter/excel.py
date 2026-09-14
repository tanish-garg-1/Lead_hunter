"""Excel workbook: one sheet per business type, each batch appended under the last with a separator row,
and the columns you edit by hand (Status, Last Contacted, Notes) synced back into memory."""
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import store

# (header, lead key, column width)
COLUMNS = [
    ("Batch", "batch_no", 7), ("Date Added", "date_added", 12), ("Business Name", "name", 30),
    ("Owner Name", "owner_name", 20), ("Phone", "phone", 18), ("Email", "email", 30),
    ("Has Website", "has_website", 12), ("Website", "website", 30), ("Website Issues", "website_issues", 40),
    ("Instagram", "instagram", 30), ("Facebook", "facebook", 30), ("WhatsApp", "whatsapp", 22),
    ("LinkedIn", "linkedin", 25), ("Address", "address", 40), ("Area", "area", 16), ("City", "city", 14),
    ("Country", "country", 12), ("Rating", "rating", 8), ("Reviews", "reviews", 9), ("Maps Link", "maps_link", 25),
    ("Lead Score", "score", 11), ("Priority", "priority", 10), ("Status", "status", 15),
    ("Last Contacted", "last_contacted", 15), ("Notes", "notes", 35), ("Place ID", "place_id", 20),
]
HEADERS = [c[0] for c in COLUMNS]
STATUSES = ["New", "Contacted", "Replied", "Meeting", "Won", "Lost", "Not interested"]
LINK_KEYS = {"website", "instagram", "facebook", "whatsapp", "linkedin", "maps_link"}
SUMMARY_SHEET = "Summary"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BATCH_FILL = PatternFill("solid", fgColor="D9E1F2")
BATCH_FONT = Font(bold=True, color="1F3864")
PRIORITY_FILLS = {
    "Hot": PatternFill("solid", fgColor="F8CBAD"),
    "Warm": PatternFill("solid", fgColor="FFE699"),
    "Cold": PatternFill("solid", fgColor="E7E6E6"),
}


class ExcelLocked(Exception):
    """The workbook is open in Excel, so it can't be saved."""


def sheet_name_for(domain):
    """'cafe' -> 'Cafes', 'grocery stores' -> 'Grocery Stores'."""
    name = " ".join(w.capitalize() for w in domain.strip().split())
    if name and not name.lower().endswith("s"):
        name += "s"
    name = re.sub(r"[\[\]:*?/\\]", "", name)[:31]
    return name if name and name != SUMMARY_SHEET else "Leads"


def _col(header):
    return HEADERS.index(header) + 1


def _create_sheet(wb, title):
    ws = wb.create_sheet(title)
    for i, (header, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, i, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.column_dimensions[get_column_letter(_col("Place ID"))].hidden = True

    status_letter = get_column_letter(_col("Status"))
    validation = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"{status_letter}2:{status_letter}20000")
    return ws


def _load(path):
    if not path.exists():
        wb = Workbook()
        wb.remove(wb.active)
        return wb
    try:
        return load_workbook(path)
    except PermissionError as exc:
        raise ExcelLocked(str(path)) from exc


def _save(wb, path):
    try:
        wb.save(path)
    except PermissionError as exc:
        raise ExcelLocked(str(path)) from exc


def append_batch(conn, path, batch):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = _load(path)
    ws = wb[batch["domain"]] if batch["domain"] in wb.sheetnames else _create_sheet(wb, batch["domain"])
    leads = store.leads_for_batch(conn, batch["id"])

    row = ws.max_row + 1
    for c in range(1, len(COLUMNS) + 1):
        ws.cell(row, c).fill = BATCH_FILL
    label = f"Batch {batch['batch_no']}  ·  {batch['created_at'][:10]}  ·  {batch['target']}  ·  {len(leads)} leads"
    ws.cell(row, 1, label).font = BATCH_FONT

    for lead in leads:
        row += 1
        for i, (_, key, _) in enumerate(COLUMNS, start=1):
            value = batch["batch_no"] if key == "batch_no" else lead[key]
            cell = ws.cell(row, i, value)
            if value and key in LINK_KEYS:
                cell.hyperlink = value if str(value).startswith("http") else f"https://{value}"
                cell.style = "Hyperlink"
            elif key == "priority" and value in PRIORITY_FILLS:
                cell.fill = PRIORITY_FILLS[value]

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{row}"
    _rebuild_summary(conn, wb)
    _save(wb, path)


def _cell_text(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def sync_from_excel(conn, path):
    """Read Status / Last Contacted / Notes from Excel into memory. Returns how many leads changed."""
    path = Path(path)
    if not path.exists():
        return 0
    wb = _load(path)
    changed = 0
    for ws in wb.worksheets:
        if ws.title == SUMMARY_SHEET:
            continue
        header_row = [c.value for c in ws[1]]
        if not all(h in header_row for h in ("Place ID", "Status", "Last Contacted", "Notes")):
            continue
        idx = {h: header_row.index(h) for h in ("Place ID", "Status", "Last Contacted", "Notes")}
        for values in ws.iter_rows(min_row=2, values_only=True):
            place_id = values[idx["Place ID"]] if len(values) > idx["Place ID"] else None
            if not place_id:
                continue  # batch separator row
            changed += store.update_tracking(
                conn, place_id,
                _cell_text(values[idx["Status"]]) or "New",
                _cell_text(values[idx["Last Contacted"]]),
                _cell_text(values[idx["Notes"]]),
            )
    conn.commit()
    return changed


def refresh_summary(conn, path):
    path = Path(path)
    if not path.exists():
        return
    wb = _load(path)
    _rebuild_summary(conn, wb)
    _save(wb, path)


def _rebuild_summary(conn, wb):
    if SUMMARY_SHEET in wb.sheetnames:
        wb.remove(wb[SUMMARY_SHEET])
    ws = wb.create_sheet(SUMMARY_SHEET, 0)
    wb.active = 0

    ws.cell(1, 1, "Lead Hunter — Summary").font = Font(bold=True, size=14)
    ws.cell(2, 1, f"Updated {datetime.now():%Y-%m-%d %H:%M}")

    headers = ["Sheet", "Total", "Hot"] + STATUSES
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(4, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(i)].width = 16
    row = 4
    for domain, entry in store.summary_rows(conn).items():
        row += 1
        values = [domain, entry["total"], entry["hot"]] + [entry["statuses"].get(s, 0) for s in STATUSES]
        for i, v in enumerate(values, start=1):
            ws.cell(row, i, v)

    row += 2
    ws.cell(row, 1, "Batches").font = Font(bold=True, size=12)
    row += 1
    for i, h in enumerate(["Sheet", "Batch", "Date", "Leads", "Target"], start=1):
        cell = ws.cell(row, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for batch in store.all_batches(conn):
        row += 1
        for i, v in enumerate(
            [batch["domain"], batch["batch_no"], batch["created_at"][:10], batch["lead_count"], batch["target"]],
            start=1,
        ):
            ws.cell(row, i, v)
