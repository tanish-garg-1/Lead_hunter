"""Offline checks (no API keys, no network): memory, dedup, Excel batches + date column, status sync,
scoring and website analysis. Run: python tests/test_offline.py"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from bs4 import BeautifulSoup
from openpyxl import load_workbook

import random

from lead_hunter import discuss, enrich, excel, filters, geo, hunt, score, store
from lead_hunter.places import search_terms

ROOT = Path(__file__).resolve().parent.parent


def fake_lead(i, **extra):
    lead = {
        "place_id": f"place-{i}", "name": f"Cafe {i}", "phone": f"+91 98765 0000{i}",
        "address": f"Shop {i}, Sector 24, Gurugram 12200{i}", "website": "", "rating": 4.4, "reviews": 80,
        "maps_link": f"https://maps.google.com/?cid={i}", "area": "Sector 24", "city": "Gurugram", "country": "India",
    }
    lead.update(extra)
    return enrich.enrich_lead(lead)  # no website -> no network


def add_batch(conn, xlsx, leads, target, day):
    batch_id, _ = store.create_batch(conn, "Cafes", target)
    for lead in leads:
        lead.update(domain="Cafes", batch_id=batch_id, date_added=day)
        lead["score"], lead["priority"] = score.score_lead(lead)
        store.insert_lead(conn, lead)
    store.finalize_batch(conn, batch_id)
    for batch in store.pending_batches(conn):
        excel.append_batch(conn, xlsx, batch)
        store.mark_exported(conn, batch["id"])


def test_batches_dedup_and_sync(tmp):
    conn = store.connect(tmp / "leads.db")
    xlsx = tmp / "leads.xlsx"
    add_batch(conn, xlsx, [fake_lead(i) for i in range(3)], "Sector 24, Gurugram, India", "2026-09-14")

    assert store.is_duplicate(conn, "place-1", "", "", "")
    assert store.is_duplicate(conn, "other-id", "098765 00001", "", ""), "same phone should be a duplicate"
    assert store.is_duplicate(conn, "other-id", "", "CAFE 2", "Somewhere 122002"), "same name + postcode"
    assert not store.is_duplicate(conn, "place-99", "+91 11111 22222", "New Cafe", "Sector 29 122009")

    add_batch(conn, xlsx, [fake_lead(i) for i in range(3, 5)], "Sector 29, Gurugram, India", "2026-09-15")

    wb = load_workbook(xlsx)
    ws = wb["Cafes"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:2] == ("Batch", "Date Added")
    assert rows[1][0].startswith("Batch 1") and "2026-" in rows[1][0]
    assert all(r[1] == "2026-09-14" for r in rows[2:5])
    assert rows[5][0].startswith("Batch 2"), "second batch goes under the first"
    assert all(r[1] == "2026-09-15" for r in rows[6:8])
    assert len(rows) == 8
    ids = [r[-1] for r in rows[1:] if r[-1]]
    assert len(ids) == len(set(ids)) == 5, "no duplicate leads in Excel"
    assert wb.sheetnames[0] == "Summary"

    place_id = ws.cell(3, len(excel.HEADERS)).value
    ws.cell(3, excel.HEADERS.index("Status") + 1, "Contacted")
    ws.cell(3, excel.HEADERS.index("Notes") + 1, "DM sent on Instagram")
    wb.save(xlsx)
    assert excel.sync_from_excel(conn, xlsx) == 1
    row = conn.execute("SELECT status, notes FROM leads WHERE place_id = ?", (place_id,)).fetchone()
    assert tuple(row) == ("Contacted", "DM sent on Instagram")
    assert excel.sync_from_excel(conn, xlsx) == 0
    assert store.summary_rows(conn)["Cafes"]["statuses"]["Contacted"] == 1
    conn.close()


def test_website_analysis():
    html = """<html><head><title>Blue Cafe</title><script src="https://static.wixstatic.com/x.js"></script></head>
    <body><a href="mailto:hello@bluecafe.in">Mail us</a>
    <a href="https://instagram.com/bluecafe/?hl=en">Instagram</a>
    <a href="https://www.facebook.com/sharer/sharer.php?u=x">Share</a>
    <a href="https://wa.me/919876500001?text=hi">WhatsApp</a>
    <p>Founded by Riya Sharma in 2015. Contact logo@2x.png</p><footer>© 2019 Blue Cafe</footer></body></html>"""
    soup = BeautifulSoup(html, "html.parser")
    text = soup.get_text(" ", strip=True)

    issues = enrich.quality_issues(soup, html, text, "http://bluecafe.in/", 5.2)
    for expected in ("No HTTPS", "Not mobile-friendly (no viewport)", "Outdated (© 2019)", "Slow (5.2s)", "Built on Wix"):
        assert expected in issues, (expected, issues)

    assert enrich.find_emails(soup, text) == {"hello@bluecafe.in"}
    socials = enrich.find_socials(soup, "http://bluecafe.in/")
    assert socials == {"instagram": "https://instagram.com/bluecafe", "whatsapp": "https://wa.me/919876500001?text=hi"}
    assert enrich.find_owner_name([text])[0] == "Riya Sharma"

    social_only = enrich.enrich_lead({"name": "X", "website": "https://www.instagram.com/xcafe"})
    assert social_only["has_website"] == "Social only" and social_only["instagram"]


def test_scoring_and_helpers():
    hot = {"has_website": "No", "rating": 4.5, "reviews": 120, "instagram": "x", "phone": "1"}
    assert score.score_lead(hot) == (80, "Hot")
    fine = {"has_website": "Yes", "website_issues": "", "rating": 3.5, "reviews": 10, "phone": "1"}
    assert score.score_lead(fine) == (10, "Cold")

    assert excel.sheet_name_for("grocery store") == "Grocery Stores"
    assert excel.sheet_name_for("cafe") == "Cafes"
    assert search_terms("cafe") == ["cafe", "coffee shop"]

    reply = 'Great choice!\n```json\n{"country": "India", "city": "Gurugram", "areas": ["Sector 24", "DLF Phase 3"], "domain": "cafe", "rationale": "busy area"}\n```'
    assert discuss.extract_target(reply)["areas"] == ["Sector 24", "DLF Phase 3"]
    assert discuss.extract_target("no json yet") is None
    bare = ('{"country": "United States", "city": "Portland", "areas": ["Pearl District", "Alberta Arts District",\n'
            '"Sellwood-Moreland"], "domain": "cafe", "rationale": "Large number of independent cafés"}')
    assert discuss.extract_target(bare)["city"] == "Portland", "JSON without the ``` fence (seen live)"
    assert discuss.extract_target("Pricing tiers: {basic} and {pro}") is None


def test_chain_and_contact_filters(tmp):
    chains = filters.load_chain_names(ROOT / "chains.txt")
    assert filters.is_known_chain("Starbucks Coffee - Sector 29", chains)
    assert filters.is_known_chain("McDonald's", chains)
    assert filters.is_known_chain("Tim Hortons (Downtown)", chains)
    assert not filters.is_known_chain("Blue Door Cafe", chains)
    assert not filters.is_known_chain("Target Fitness Studio", chains)
    assert not filters.is_known_chain("Calli. The coffee club", chains), "brand must be at the start of the name"

    detector = filters.ChainDetector(set())
    page = [
        {"place_id": "a", "name": "Brew Co - Downtown", "website": "https://brewco.com/locations/downtown"},
        {"place_id": "b", "name": "Brew Co (Uptown)", "website": ""},
        {"place_id": "c", "name": "Solo Cafe", "website": "https://instagram.com/solo"},
        {"place_id": "d", "name": "Other Cafe", "website": "https://instagram.com/other"},
    ]
    keys = detector.observe(page)
    assert detector.reason(page[0], keys) and detector.reason(page[1], keys), "two branches on one page"
    assert not detector.reason(page[2], keys) and not detector.reason(page[3], keys), "shared instagram.com isn't a chain"
    assert detector.reason({"place_id": "z", "name": "Local Mart", "brand": True}, {})

    repeat = filters.ChainDetector(set())
    solo = {"place_id": "x", "name": "Solo", "website": "https://solo.cafe"}
    for _ in range(3):  # same café returned by three different queries
        keys = repeat.observe([solo])
    assert not repeat.reason(solo, keys) and not repeat.has_many_branches(solo)

    assert not filters.has_contact({"name": "x", "website": "https://x.com"}), "website alone isn't a contact"
    assert filters.has_contact({"instagram": "https://instagram.com/x"})
    assert filters.has_contact({"phone": "+44 20 1234 5678"})

    conn = store.connect(tmp / "skip.db")
    store.add_skipped(conn, "p1", "Quiet Cafe", "no contact details")
    assert store.is_skipped(conn, "p1")
    conn.execute("UPDATE skipped SET skipped_on = '2020-01-01'")
    assert not store.is_skipped(conn, "p1"), "re-check after 90 days"
    conn.close()


def test_geo_helpers():
    square = [[(13.0, 52.0), (14.0, 52.0), (14.0, 53.0), (13.0, 53.0), (13.0, 52.0)]]
    assert geo.inside_city(square, 13.5, 52.5)
    assert not geo.inside_city(square, 15.0, 52.5)
    assert geo.box_touches_city((52.9, 13.9, 53.1, 14.1), square), "corner overlap counts"
    assert not geo.box_touches_city((54.0, 15.0, 54.1, 15.1), square)
    box = (52.0, 13.0, 52.2, 13.2)
    parts = geo.split_box(box)
    assert len(parts) == 4 and abs(sum(geo.box_area_km2(p) for p in parts) / geo.box_area_km2(box) - 1) < 0.01
    assert len(geo.initial_grid((52.3, 13.0, 52.7, 13.8), 4.0)) > 50


class FakeMap:
    """Stands in for Google: returns places inside a square, max 20 per 'page', saturated when more exist."""
    name = "fake"

    def __init__(self, places):
        self.places = places
        self.calls = 0

    def search_box(self, domain, box, can_split):
        self.calls += 1
        south, west, north, east = box
        inside = [dict(p) for p in self.places if south <= p["lat"] < north and west <= p["lng"] < east]
        if len(inside) > 20 and can_split:
            return inside[:20], True
        return inside[:60], False


def make_city(rng):
    places = []
    for i in range(160):  # dense centre, two postcodes
        lat, lng = 52.52 + rng.uniform(-0.012, 0.012), 13.40 + rng.uniform(-0.02, 0.02)
        places.append({"place_id": f"c{i}", "name": f"Cafe {i}", "phone": f"+49 30 5550{i:04d}",
                       "address": f"Street {i}", "lat": lat, "lng": lng,
                       "postcode": "10115" if lat > 52.52 else "10997", "suburb": "Mitte" if lat > 52.52 else "Kreuzberg"})
    for i in range(40):  # outskirts
        lat, lng = 52.52 + rng.uniform(-0.045, 0.045), 13.40 + rng.uniform(-0.075, 0.075)
        places.append({"place_id": f"o{i}", "name": f"Edge Cafe {i}", "phone": f"+49 30 6660{i:04d}",
                       "address": f"Road {i}", "lat": lat, "lng": lng, "postcode": "13000", "suburb": "Outer"})
    for i in range(4):  # a local chain spread over the city
        places.append({"place_id": f"k{i}", "name": f"Bean Chain - Branch {i}", "phone": f"+49 30 7770{i:04d}",
                       "address": f"Chain {i}", "lat": 52.49 + i * 0.02, "lng": 13.34 + i * 0.04,
                       "postcode": "12000", "suburb": "Various"})
    return places


def test_city_hunt(tmp):
    conn = store.connect(tmp / "hunt.db")
    city_geo = {"name": "Testville", "lat": 52.52, "lng": 13.40, "bbox": (52.47, 13.32, 52.57, 13.48), "polygons": []}
    hunt_id = hunt.create_hunt(conn, "cafe", "Testville", "Germany", city_geo)
    fake = FakeMap(make_city(random.Random(7)))

    # Day 1: ask for 30. Central squares go first, and the leads come in one cluster.
    def no_lookup(lat, lng):
        raise AssertionError("places with a suburb shouldn't need a name lookup")

    def open_hunt(hid):
        return hunt.Hunt(conn, hid, fake, filters.ChainDetector(set()), log=lambda *_: None, namer=no_lookup)

    day1 = open_hunt(hunt_id)
    day1.fill_queue(30)
    picked = day1.take(30)
    assert len(picked) == 30
    order = [p["cluster"] for p in picked]
    switches = sum(a != b for a, b in zip(order, order[1:]))
    assert switches == len(set(order)) - 1 <= 1, "leads come grouped by neighbourhood, one cluster at a time"
    assert picked[0]["cluster"] in ("Mitte", "Kreuzberg"), "central neighbourhoods first"
    for p in picked:
        assert not day1.recheck(p, claim=True)
        p.update(domain="Cafes", batch_id=1, website="")
        store.insert_lead(conn, p)
    conn.commit()
    day1.done([p["place_id"] for p in picked])
    calls_day1 = fake.calls
    assert hunt.progress(conn, hunt_id)["covered_pct"] < 100

    # Day 2+: a fresh session resumes; keep going until the whole city is covered.
    day2 = open_hunt(hunt_id)
    assert calls_day1 > 0
    saved = {p["place_id"] for p in picked}
    while True:
        day2.fill_queue(50)
        batch = day2.take(50)
        if not batch:
            break
        for p in batch:
            if not day2.recheck(p, claim=True):
                saved.add(p["place_id"])
                p.update(domain="Cafes", batch_id=1)
                store.insert_lead(conn, p)
        conn.commit()
        day2.done([p["place_id"] for p in batch])

    assert day2.is_complete()
    assert hunt.progress(conn, hunt_id)["covered_pct"] == 100
    independents = {p["place_id"] for p in fake.places if not p["place_id"].startswith("k")}
    assert independents <= saved, f"missed: {sorted(independents - saved)[:5]}"
    chains_saved = {pid for pid in saved if pid.startswith("k")}
    assert len(chains_saved) <= 2, "a chain is caught once its 3rd branch shows up"
    count = conn.execute("SELECT COUNT(*) FROM leads").fetchone()[0]
    assert count == len(saved), "no duplicates"
    assert fake.calls < 120, f"too many searches: {fake.calls}"

    conn.close()

    # Focus: in a fresh database (so nothing is 'already in your list'), a chosen neighbourhood is served first.
    conn = store.connect(tmp / "focus.db")
    hunt_id2 = hunt.create_hunt(conn, "cafe", "Testville", "Germany", city_geo)
    other = open_hunt(hunt_id2)
    other.fill_queue(200)
    hunt.set_focus(conn, hunt_id2, "Outer")
    focused = open_hunt(hunt_id2)
    assert all(p["cluster"] == "Outer" for p in focused.take(5))
    names = {c["cluster"] for c in hunt.clusters(conn, hunt_id2)}
    assert {"Mitte", "Kreuzberg", "Outer"} <= names

    # A place without a suburb gets its neighbourhood looked up once per ~1 km zone, then cached.
    lookups = []
    namer_hunt = hunt.Hunt(conn, hunt_id2, fake, filters.ChainDetector(set()), log=lambda *_: None,
                           namer=lambda lat, lng: lookups.append(1) or "Prenzlauer Berg")
    assert namer_hunt._neighbourhood(52.5401, 13.4101) == "Prenzlauer Berg"
    assert namer_hunt._neighbourhood(52.5399, 13.4099) == "Prenzlauer Berg" and len(lookups) == 1
    assert hunt.cluster_of({"lat": 52.54, "lng": 13.41}, 5, namer_hunt._neighbourhood) == ("Prenzlauer Berg", "Prenzlauer Berg")
    conn.close()
    return fake.calls


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as tmp:
        test_batches_dedup_and_sync(Path(tmp))
        test_chain_and_contact_filters(Path(tmp))
        test_geo_helpers()
        calls = test_city_hunt(Path(tmp))
        print(f"City hunt covered every independent cafe in the fake city using {calls} map searches")
    test_website_analysis()
    test_scoring_and_helpers()
    print("All offline tests passed")
